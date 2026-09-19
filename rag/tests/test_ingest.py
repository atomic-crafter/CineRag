import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import vectorstore
from chunker import chunk_document
from loaders import load_document


def test_load_pdf(tmp_path):
    pdf_path = _make_sample_pdf(tmp_path)
    pages = load_document(str(pdf_path))
    assert len(pages) == 1
    page_number, text = pages[0]
    assert page_number == 1
    assert "Hello RAG world" in text


def test_load_pptx(tmp_path):
    pptx_path = _make_sample_pptx(tmp_path)
    slides = load_document(str(pptx_path))
    assert len(slides) == 1
    slide_number, text = slides[0]
    assert slide_number == 1
    assert "Hello slide" in text


def test_load_unsupported_extension(tmp_path):
    bad_path = tmp_path / "notes.txt"
    bad_path.write_text("hi")
    try:
        load_document(str(bad_path))
        assert False, "expected ValueError"
    except ValueError:
        pass


def test_chunk_document_preserves_page_number_and_indexes():
    long_text = "word " * 2000  # long enough to force multiple chunks
    pages = [(1, long_text), (2, "short page text")]
    chunks = chunk_document(pages)

    assert len(chunks) > 2
    # page 1 chunks come first, indices are sequential from 0
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    page1_chunks = [c for c in chunks if c.page_number == 1]
    page2_chunks = [c for c in chunks if c.page_number == 2]
    assert len(page1_chunks) > 1
    assert len(page2_chunks) == 1
    assert page2_chunks[0].chunk_text.strip() == "short page text"


def test_chunk_document_empty_pages_produce_no_chunks():
    assert chunk_document([]) == []
    assert chunk_document([(1, "   ")]) == []



def test_ingest_query_delete_roundtrip(client, tmp_path):
    file_id = "test-ingest-roundtrip"
    vectorstore.delete_document(file_id)  # clean slate in case of leftovers

    pdf_path = _make_sample_pdf(tmp_path)
    resp = client.post(
        "/ingest",
        json={"file_id": file_id, "filename": "sample.pdf", "file_path": str(pdf_path)},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body == {"status": "success", "file_id": file_id, "chunks_added": 1}

    # re-ingesting the same file_id replaces rather than duplicates chunks
    resp = client.post(
        "/ingest",
        json={"file_id": file_id, "filename": "sample.pdf", "file_path": str(pdf_path)},
    )
    assert resp.json()["chunks_added"] == 1

    resp = client.delete(f"/ingest/{file_id}")
    assert resp.status_code == 200
    assert resp.json() == {"status": "deleted", "file_id": file_id, "chunks_removed": 1}


def test_ingest_missing_file_returns_404(client):
    resp = client.post(
        "/ingest",
        json={"file_id": "missing", "filename": "x.pdf", "file_path": "/no/such/file.pdf"},
    )
    assert resp.status_code == 404


def _make_sample_pdf(tmp_path) -> Path:
    pdf_path = tmp_path / "sample.pdf"
    text = "Hello RAG world"
    content = f"BT /F1 24 Tf 72 100 Td ({text}) Tj ET".encode()

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 5 0 R >> >> "
        b"/MediaBox [0 0 200 200] /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    buf = bytearray(b"%PDF-1.4\n")
    offsets = [0]  # object 0 is the free-list head
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(buf))
        buf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_offset = len(buf)
    n = len(objects) + 1
    buf += f"xref\n0 {n}\n".encode()
    buf += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        buf += f"{off:010d} 00000 n \n".encode()
    buf += (
        f"trailer\n<< /Size {n} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF"
    ).encode()

    pdf_path.write_bytes(bytes(buf))
    return pdf_path


def _make_sample_pptx(tmp_path) -> Path:
    from pptx import Presentation

    pptx_path = tmp_path / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 100, 100)
    textbox.text_frame.text = "Hello slide"
    prs.save(str(pptx_path))
    return pptx_path
